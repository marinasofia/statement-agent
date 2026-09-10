import json
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from core.review import ReviewQueue
from review import main


@pytest.fixture
def item(tmp_path):
    queue = ReviewQueue(tmp_path / "review.sqlite")
    data = {
        "account_holder": "=SYNTHETIC",
        "closing_balance": 110,
        "opening_balance": 100,
        "statement_date": "2026-08-31",
        "statement_period_start": "2026-08-01",
        "statement_period_end": "2026-08-31",
        "transactions": [
            {"date": "2026-08-02", "description": "Synthetic credit", "amount": 10}
        ],
    }
    identity = queue.add(
        {
            "validated_data": data,
            "raw_text": "Synthetic statement",
            "file_path": "synthetic.pdf",
        },
        b"synthetic pdf",
    )
    return queue, identity


def test_approval_required_and_duplicates_preserve_corrections(item, tmp_path):
    queue, identity = item
    with pytest.raises(ValueError, match="Only approved"):
        queue.export(identity, tmp_path)
    before = queue.get(identity)
    changed = {**before["current"], "account_holder": "Corrected name"}
    queue.change(identity, 0, "Reviewer", "Source name checked", data=changed)
    assert (
        queue.add({"validated_data": before["original"]}, b"synthetic pdf") == identity
    )
    recovered = ReviewQueue(queue.path).get(identity)
    assert recovered["current"]["account_holder"] == "Corrected name"
    assert recovered["original"]["account_holder"] == "=SYNTHETIC"
    with pytest.raises(ValueError, match="reload"):
        queue.change(identity, 0, "Reviewer", "stale approval", approve=True)
    queue.change(
        identity, 1, "Reviewer", "Compared all fields with source", approve=True
    )
    path = queue.export(identity, tmp_path / "export")
    assert queue.export(identity, tmp_path / "export") == path
    wb = load_workbook(path)
    assert wb["2026-08"].max_row == 2
    wb.close()
    queue.change(identity, 2, "Reviewer", "Another correction", data=changed)
    assert queue.get(identity)["status"] == "NEEDS_REVIEW"


def test_skipped_and_failed_checks_block_silent_approval(item):
    queue, identity = item
    data = queue.get(identity)["current"]
    data["opening_balance"] = None
    queue.change(identity, 0, "Reviewer", "Not present on source", data=data)
    with pytest.raises(ValueError, match="Skipped"):
        queue.change(identity, 1, "Reviewer", "Reviewed", approve=True)
    queue.change(
        identity,
        1,
        "Reviewer",
        "Source has no opening; independently checked rows",
        approve=True,
        acknowledge_skipped=True,
    )
    data["opening_balance"] = 500
    queue.change(identity, 2, "Reviewer", "Correction", data=data)
    with pytest.raises(ValueError, match="failed checks"):
        queue.change(
            identity, 3, "Reviewer", "Approve", approve=True, acknowledge_skipped=True
        )


def test_export_failure_survives_restart_and_can_retry(item, tmp_path):
    queue, identity = item
    queue.change(identity, 0, "Reviewer", "Compared to source", approve=True)
    with patch("core.review.write_workbook", side_effect=OSError("disk full")):
        with pytest.raises(OSError):
            queue.export(identity, tmp_path / "exports")
    restarted = ReviewQueue(queue.path)
    assert restarted.get(identity)["status"] == "APPROVED"
    assert restarted.get(identity)["history"][-1]["action"] == "EXPORT_FAILED"
    assert restarted.export(identity, tmp_path / "exports").exists()


def test_invalid_correction_rolls_back_and_cli_reads(item, capsys):
    queue, identity = item
    with pytest.raises(ValueError):
        queue.change(
            identity, 0, "Reviewer", "Invalid", data={"closing_balance": float("nan")}
        )
    assert queue.get(identity)["revision"] == 0
    assert main(["--db", str(queue.path), "show", identity]) == 0
    assert (
        json.loads(capsys.readouterr().out)["source"]["text"] == "Synthetic statement"
    )


def test_failed_extraction_is_not_approvable(tmp_path):
    queue = ReviewQueue(tmp_path / "queue.sqlite")
    identity = queue.add({"error": "Image only"}, b"image pdf")
    with pytest.raises(ValueError, match="Failed extraction"):
        queue.change(identity, 0, "Reviewer", "Approve", approve=True)


def test_failed_extraction_can_be_retried_with_audit(item):
    queue, identity = item
    failed = queue.add({"error": "Temporary extraction failure"}, b"retry pdf")
    queue.add({"validated_data": queue.get(identity)["current"]}, b"retry pdf")
    recovered = queue.get(failed)
    assert recovered["status"] == "NEEDS_REVIEW"
    assert recovered["history"][-1]["action"] == "REEXTRACTED"


def test_preview_escapes_untrusted_content(item, tmp_path):
    queue, identity = item
    data = queue.get(identity)["current"]
    data["account_holder"] = "<script>alert(1)</script>"
    queue.change(identity, 0, "Reviewer", "Synthetic malicious field", data=data)
    output = queue.preview(identity, tmp_path / "preview.html").read_text()
    assert "<script>" not in output
    assert "&lt;script&gt;" in output


def test_batch_resume_skips_saved_items_and_does_not_export(
    item, tmp_path, monkeypatch
):
    from unittest.mock import Mock

    import run_batch

    queue, identity = item
    folder = tmp_path / "inputs"
    folder.mkdir()
    source = folder / "synthetic.pdf"
    source.write_bytes(b"synthetic pdf")
    monkeypatch.setattr(run_batch, "ALLOWED_UPLOAD_DIR", str(folder))
    processor = Mock(
        return_value={
            "validated_data": queue.get(identity)["current"],
            "file_path": str(source),
        }
    )
    writer = Mock()
    monkeypatch.setattr(run_batch, "process_pdf", processor)
    monkeypatch.setattr(run_batch, "write_workbook", writer)
    assert run_batch.main(["--input", str(folder), "--review-db", str(queue.path)]) == 0
    processor.assert_not_called()
    writer.assert_not_called()
    source.write_bytes(b"new synthetic pdf")
    assert run_batch.main(["--input", str(folder), "--review-db", str(queue.path)]) == 0
    assert processor.call_count == 1
    assert len(queue.list()) == 2
    writer.assert_not_called()
