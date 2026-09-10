"""Local, single-operator review queue. SQLite commits preserve review decisions."""

import hashlib
import html
import json
import os
import sqlite3
import tempfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from agents.statement_extraction.reconcile import reconcile
from agents.statement_extraction.schema import StatementData
from core.excel import _append_literal_row, write_workbook


def encoded(value):
    return json.dumps(value, sort_keys=True, allow_nan=False)


class ReviewQueue:
    def __init__(self, path):
        self.path = Path(path)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.connection() as db:
            db.execute(
                "CREATE TABLE IF NOT EXISTS items (id TEXT PRIMARY KEY, original TEXT NOT NULL, current TEXT NOT NULL, source TEXT NOT NULL, status TEXT NOT NULL, revision INTEGER NOT NULL)"
            )
            db.execute(
                "CREATE TABLE IF NOT EXISTS events (seq INTEGER PRIMARY KEY, item TEXT NOT NULL, time TEXT NOT NULL, action TEXT NOT NULL, detail TEXT NOT NULL)"
            )

    @contextmanager
    def connection(self):
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        try:
            with db:
                yield db
        finally:
            db.close()

    def event(self, db, item, action, detail):
        db.execute(
            "INSERT INTO events(item,time,action,detail) VALUES(?,?,?,?)",
            (item, datetime.now(timezone.utc).isoformat(), action, encoded(detail)),
        )

    def add(self, state, source_bytes):
        identity = hashlib.sha256(source_bytes).hexdigest()
        data = state.get("validated_data")
        if data is not None:
            data = StatementData.model_validate(data).model_dump()
        status = "FAILED" if state.get("error") or data is None else "NEEDS_REVIEW"
        source = {
            "file": Path(state.get("file_path", "")).name,
            "text": state.get("raw_text", ""),
            "error": state.get("error"),
        }
        with self.connection() as db:
            inserted = db.execute(
                "INSERT OR IGNORE INTO items VALUES(?,?,?,?,?,0)",
                (identity, encoded(data), encoded(data), encoded(source), status),
            ).rowcount
            if not inserted and status != "FAILED":
                previous = db.execute(
                    "SELECT * FROM items WHERE id=?", (identity,)
                ).fetchone()
                if previous["status"] == "FAILED":
                    self.event(
                        db,
                        identity,
                        "REEXTRACTED",
                        {
                            "previous_source": json.loads(previous["source"]),
                            "data": data,
                        },
                    )
                    db.execute(
                        "UPDATE items SET original=?,current=?,source=?,status=?,revision=revision+1 WHERE id=?",
                        (
                            encoded(data),
                            encoded(data),
                            encoded(source),
                            status,
                            identity,
                        ),
                    )
            if inserted:
                self.event(db, identity, "IMPORTED", {"status": status})
        return identity

    def list(self):
        with self.connection() as db:
            return [
                dict(r)
                for r in db.execute(
                    "SELECT id,status,revision FROM items ORDER BY rowid"
                )
            ]

    def get(self, identity):
        with self.connection() as db:
            row = db.execute("SELECT * FROM items WHERE id=?", (identity,)).fetchone()
            if row is None:
                raise ValueError("Unknown review item")
            result = dict(row)
            for key in ("original", "current", "source"):
                result[key] = json.loads(result[key])
            result["checks"] = (
                reconcile(result["current"]).as_dict() if result["current"] else None
            )
            result["history"] = [
                {**dict(e), "detail": json.loads(e["detail"])}
                for e in db.execute(
                    "SELECT time,action,detail FROM events WHERE item=? ORDER BY seq",
                    (identity,),
                )
            ]
            return result

    def change(
        self,
        identity,
        revision,
        actor,
        reason,
        data=None,
        approve=False,
        acknowledge_skipped=False,
    ):
        if not actor.strip() or not reason.strip():
            raise ValueError("Reviewer and reason are required")
        with self.connection() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT * FROM items WHERE id=?", (identity,)).fetchone()
            if row is None or row["revision"] != revision:
                raise ValueError(
                    "Item changed or does not exist; reload before editing"
                )
            current = json.loads(row["current"])
            if approve:
                if current is None or row["status"] == "FAILED":
                    raise ValueError("Failed extraction cannot be approved")
                checks = reconcile(current)
                if not checks.ok:
                    raise ValueError("Correct failed checks before approval")
                if (
                    any(c.outcome == "skipped" for c in checks.checks)
                    and not acknowledge_skipped
                ):
                    raise ValueError("Skipped checks require explicit acknowledgement")
                status, action = "APPROVED", "APPROVED"
            else:
                if current is None:
                    raise ValueError(
                        "Failed extraction must be re-extracted before correction"
                    )
                if not isinstance(data, dict) or set(data) - set(
                    StatementData.model_fields
                ):
                    raise ValueError("Correction contains unknown statement fields")
                for transaction in data.get("transactions") or []:
                    if not isinstance(transaction, dict) or set(transaction) - {
                        "date",
                        "description",
                        "amount",
                    }:
                        raise ValueError(
                            "Correction contains unknown transaction fields"
                        )
                current = StatementData.model_validate(data).model_dump()
                status, action = "NEEDS_REVIEW", "CORRECTED"
            self.event(
                db,
                identity,
                action,
                {
                    "actor": actor,
                    "reason": reason,
                    "revision": revision + 1,
                    "before": json.loads(row["current"]),
                    "after": current,
                    "acknowledged_skipped": acknowledge_skipped if approve else False,
                },
            )
            db.execute(
                "UPDATE items SET current=?,status=?,revision=revision+1 WHERE id=?",
                (encoded(current), status, identity),
            )

    def export(self, identity, directory):
        # Hold the writer lock through export so edits cannot race an approval.
        with self.connection() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT * FROM items WHERE id=?", (identity,)).fetchone()
            if row is None or row["status"] != "APPROVED":
                raise ValueError("Only approved items can be exported")
            data = json.loads(row["current"])
            target = Path(directory) / f"{identity}-r{row['revision']}.xlsx"
            try:
                target.parent.mkdir(parents=True, exist_ok=True)
                with tempfile.TemporaryDirectory(dir=target.parent) as temporary:
                    staging = Path(temporary) / "approved.xlsx"
                    write_workbook(
                        [
                            {
                                "validated_data": data,
                                "status": "OK",
                                "statement_month": (data.get("statement_date") or "")[
                                    :7
                                ]
                                or None,
                            }
                        ],
                        staging,
                        "default",
                    )
                    workbook = load_workbook(staging)
                    try:
                        transactions = workbook.create_sheet("Transactions")
                        _append_literal_row(
                            transactions, ["Date", "Description", "Amount"]
                        )
                        for transaction in data.get("transactions") or []:
                            _append_literal_row(
                                transactions,
                                [
                                    transaction.get("date"),
                                    transaction["description"],
                                    transaction["amount"],
                                ],
                            )
                        audit = workbook.create_sheet("Approval history")
                        _append_literal_row(
                            audit,
                            ["Source SHA256", "Revision", "Time", "Action", "Detail"],
                        )
                        for event in db.execute(
                            "SELECT * FROM events WHERE item=? ORDER BY seq",
                            (identity,),
                        ):
                            if event["action"] in ("IMPORTED", "CORRECTED", "APPROVED"):
                                _append_literal_row(
                                    audit,
                                    [
                                        identity,
                                        row["revision"],
                                        event["time"],
                                        event["action"],
                                        event["detail"],
                                    ],
                                )
                        workbook.save(staging)
                    finally:
                        workbook.close()
                    os.replace(staging, target)
            except Exception as exc:
                self.event(
                    db,
                    identity,
                    "EXPORT_FAILED",
                    {"revision": row["revision"], "error_type": type(exc).__name__},
                )
                failure = exc
            else:
                self.event(
                    db,
                    identity,
                    "EXPORTED",
                    {"revision": row["revision"], "file": target.name},
                )
                failure = None
        if failure is not None:
            raise failure
        return target

    def preview(self, identity, output):
        item = self.get(identity)
        def escape(value):
            return html.escape(str(value))
        sections = []
        for title, value in [
            ("Original extraction", item["original"]),
            ("Current values", item["current"]),
            ("Automated checks", item["checks"]),
            ("Review history", item["history"]),
        ]:
            sections.append(
                f"<section><h2>{title}</h2><pre>{escape(json.dumps(value, indent=2))}</pre></section>"
            )
        source = escape(
            item["source"].get("text")
            or "Source text unavailable. Compare against the original PDF before approval."
        )
        document = '<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>Statement review</title><style>body{font:16px system-ui;margin:32px;max-width:1200px;color:#202020;background:#fff}main{display:grid;grid-template-columns:1fr 1fr;gap:24px}pre{white-space:pre-wrap;overflow-wrap:anywhere;font-size:13px;padding:16px;background:#f3f3f3}h1{font-size:24px}h2{font-size:18px}@media(max-width:700px){main{grid-template-columns:1fr}}</style>'
        document += f"<h1>Statement review</h1><p>Status: {escape(item['status'])}. Revision: {item['revision']}. This is a read-only snapshot.</p><p>Source: {escape(item['source'].get('file'))}. Automated checks do not establish factual accuracy.</p><main><section><h2>Extracted source text</h2><pre>{source}</pre></section><div>{''.join(sections)}</div></main></html>"
        Path(output).write_text(document, encoding="utf-8")
        return Path(output)
